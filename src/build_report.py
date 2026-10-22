"""Generate the laboratory report from measured result files."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DATA, RESULTS, LIT = ROOT / "data", ROOT / "results", ROOT / "literature"
OUTPUT = ROOT / "Лабораторная_2_Мультимодальная_адаптация_домена.docx"


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


schemes = read_csv(RESULTS / "scheme_intervals.csv")
matched = read_csv(RESULTS / "paired_controls.csv")
chosen = read_csv(RESULTS / "paired_comparisons.csv")
ablations = read_csv(RESULTS / "missing_modality.csv")
retrained = read_csv(RESULTS / "retrained_ablation_summary.csv")
errors = read_csv(RESULTS / "test_errors.csv")
literature = read_csv(LIT / "matrix.csv")
audit = json.loads((DATA / "audit.json").read_text())
selection = json.loads((RESULTS / "selection.json").read_text())
extract = {k: json.loads((DATA / f"{k}_extraction.json").read_text())
           for k in ("officehome", "digits")}
manifest = json.loads((ROOT.parent / "lab1_officehome/data/manifest.json").read_text())
class_names = {int(row["label"]): row["class_name"].replace("_", " ") for row in manifest}


def lookup(rows, **terms):
    return next(r for r in rows if all(r.get(k) == v for k, v in terms.items()))


def fmt(x, n=3):
    return f"{float(x):.{n}f}".replace(".", ",")


def ci(row):
    return f"[{fmt(row['ci_low'])}; {fmt(row['ci_high'])}]"


def score_ci(row):
    return f"[{fmt(row['macro_f1_ci_low'])}; {fmt(row['macro_f1_ci_high'])}]"


doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(2), Cm(1.8)
section.left_margin, section.right_margin = Cm(2), Cm(1.8)

for name, size, bold, before, after in [
    ("Normal", 11, False, 0, 5),
    ("Title", 16, True, 0, 10),
    ("Heading 1", 13, True, 12, 5),
    ("Heading 2", 11.5, True, 8, 4),
    ("Caption", 9.5, False, 5, 5),
]:
    style = doc.styles[name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.15

# The built-in Word Title style carries a colored bottom rule on some systems.
title_properties = doc.styles["Title"]._element.get_or_add_pPr()
title_border = OxmlElement("w:pBdr")
title_bottom = OxmlElement("w:bottom")
title_bottom.set(qn("w:val"), "nil")
title_border.append(title_bottom)
title_properties.append(title_border)


def paragraph(text="", style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    p.paragraph_format.keep_together = True
    return p


def heading(text, level=1):
    return doc.add_heading(text, level)


def caption(text):
    p = paragraph(text, "Caption")
    p.paragraph_format.keep_with_next = True
    return p


def table(headers, rows, widths=None, font=8.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = str(h)
    for data in rows:
        cells = t.add_row().cells
        for i, value in enumerate(data):
            cells[i].text = str(value)
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Cm(widths[ci])
            tcpr = cell._tc.get_or_add_tcPr()
            if ri == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "DDE5EB")
                tcpr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(font)
                    run.font.bold = ri == 0
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def hyperlink(p, label, url):
    rel = p.part.relate_to(url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    node = OxmlElement("w:hyperlink")
    node.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1A4D7A")
    props.append(color)
    run.append(props)
    txt = OxmlElement("w:t")
    txt.text = label
    run.append(txt)
    node.append(run)
    p._p.append(node)


footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

for line in ("Финансовый университет при Правительстве Российской Федерации",
             "Департамент анализа данных и машинного обучения"):
    p = paragraph(line)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
doc.add_paragraph()
p = paragraph("Лабораторная работа 2\nМультимодальные модели для адаптации визуального домена", "Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = paragraph("По теме ВКР «Сравнительный анализ методов выравнивания распределений для адаптации домена в задачах компьютерного зрения»")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()
p = paragraph("Выполнила: А. Д. Якумова")
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p = paragraph("2026")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

heading("Аннотация")
paragraph("Проверено, помогают ли текстовые описания и числовые характеристики изображения перенести распознавание объектов из фотографий товаров Product на рисунки Clipart. Изображение и текст кодировались предобученной SigLIP 2; 19 числовых характеристик измерялись непосредственно на исходном изображении. Сравнены восемь схем объединения признаков, MMD и CORAL, одномодальные контроли и zero-shot распознавание. На Clipart лучшим ориентиром оказалось прямое zero-shot применение SigLIP 2: Macro-F1 0,879 против 0,634 у выбранной по Product схемы объединения. Для Office-Home добавление зависимых от изображения текста и чисел не улучшило перенос. На дополнительном наборе Digits обнаружен эффект выравнивания, но его нельзя приписать объединению без сравнения с визуальной моделью с тем же выравниванием.")

heading("1 Постановка и связь с темой ВКР")
paragraph("Адаптация домена изучает ситуацию, когда обучающие и рабочие изображения имеют разные распределения. Здесь метки объектов доступны для фотографий Product, а целевые рисунки Clipart при настройке выравнивания считаются неразмеченными. MMD сближает распределения признаков через расстояние между средними в ядровом пространстве; CORAL сближает ковариационные матрицы. Оба метода не требуют истинных меток Clipart. Основная метрика — Macro-F1, то есть среднее F1 по всем 65 классам с одинаковым весом каждого класса.")
paragraph("SigLIP 2 — предобученный кодировщик изображения и текста с общей семантической областью признаков. Это мультимодальная модель, но не генеративная LLM: в данном опыте она не пишет ответ, а возвращает числовые векторы и сходство изображения с текстовым названием класса. Текстовый вход для схем объединения — первые шесть автоматических меток Apple Vision для того же изображения, без обращения к истинному классу. Числовой вход — статистики того же изображения. Поэтому работа исследует полезность трёх представлений одного объекта, а не независимых сенсоров.")

heading("2 Обзор исследований и исследовательский вопрос")
paragraph("Обзор охватывает 15 первичных работ. Office-Home задаёт визуальные домены; DAN и Deep CORAL — базовые способы выравнивания; CLIP, SigLIP 2 и DPA показывают, как семантика текста может участвовать в распознавании и адаптации. Работы по тексту, таблицам, временным рядам и финансам нужны для сопоставления с полным заданием лабораторной. Они не служат доказательством, что Office-Home содержит временной ряд. Полная матрица с кодированием, схемой объединения, данными, метриками, доступностью кода и ограничениями сохранена в literature/matrix.csv.")
caption("Таблица 1 — Сопоставление 15 работ по способу объединения и применимости")
lit_rows = [[r["id"], f"{r['authors']} ({r['year']})", r["fusion"],
             r["task_data"], "ссылка в матрице" if r["code"].startswith("http") else "не проверен"]
            for r in literature]
table(["№", "Работа", "Выравнивание или объединение", "Задача и данные", "Код"],
      lit_rows, [0.7, 2.9, 5.0, 4.1, 3.0], 7.7)
paragraph("Исследовательский вопрос: сохраняется ли польза MMD/CORAL и добавочных представлений, если исходный визуальный вектор уже получен сильной vision-language моделью? Для ответа нужна проверка не только против слабого визуального классификатора, но и против zero-shot SigLIP 2 и визуальной модели с тем же методом выравнивания.")

heading("3 Данные и схема мультимодального объекта")
paragraph("Объект объединяется строго по паре dataset + id. Для Office-Home это JPEG, текст автоматических категорий Apple Vision и 19 числовых величин: средние и стандартные отклонения каналов RGB, насыщенности и яркости, характеристики серого изображения, доли тёмных и светлых пикселей, величина и доля границ, логарифм отношения ширины к высоте. Они рассчитаны один раз на исходном изображении без изменения цвета или яркости. Метка класса хранится отдельно и не поступает на вход кодировщика.")
caption("Таблица 2 — Паспорт двух наборов и назначение частей")
table(["Набор", "Часть", "Домен", "N", "Назначение"], [
    ["Office-Home", "source_train", "Product", 520, "обучение головы"],
    ["Office-Home", "validation", "Product", 130, "выбор эпохи и схемы"],
    ["Office-Home", "target_adapt", "Clipart", 195, "безметочное выравнивание"],
    ["Office-Home", "test", "Clipart", 130, "основная оценка"],
    ["Office-Home", "new_slice / external", "Art / Real World", "130 / 130", "другие домены"],
    ["Digits", "source_train / validation", "исходный", "790 / 198", "обучение / выбор"],
    ["Digits", "target_adapt / test", "инверсия и окраска", "404 / 405", "адаптация / оценка"],
], [2.4, 3.0, 3.0, 1.7, 5.6], 8.3)
paragraph("Office-Home содержит 1235 объектов 65 классов; на основном тесте только по два рисунка на класс. Digits содержит 1797 разных рукописных цифр десяти классов: целевой стиль сформирован инверсией, окраской и слабым размытием уже выделенных объектов. Тексты Digits — детерминированное описание геометрии штрихов без названия цифры. Пропусков текста, нечисловых признаков и пересечений групп между частями аудит не обнаружил; точных повторов исходных Digits тоже нет. Результаты аудита и исключений: data/audit.json и data/exclusions.csv.")
paragraph("Office-Home и Digits не содержат наблюдений по календарному времени. Связь по id является объектной синхронизацией; UTC, окно до момента прогноза и отложенный временной период здесь неприменимы. Следовательно, требование задания о настоящем числовом временном ряде и проверке на будущем периоде этим выбором данных не выполнено. Подменять его искусственными стадиями обработки изображения было бы методически неверно.")

heading("4 Модель, схемы и протокол")
paragraph("Использован официальный checkpoint google/siglip2-base-patch16-224, ревизия 75de2d55ec2d0b4efc50b3e9ad70dba96a7b2fa2. Замороженные кодировщики дали для каждого объекта 768-мерные векторы изображения и текста. Zero-shot вариант сравнивает вектор изображения с заранее заданными текстами 65 классов Office-Home или десяти цифр и не обучается на целевых рисунках. Три входа классификационной головы — визуальный вектор, текстовый вектор и 19 чисел. У каждой ветви линейное отображение в 128 измерений и ReLU; затем общее отображение и выход по классам.")
caption("Таблица 3 — Восемь проверенных схем объединения")
table(["Схема", "Объединение", "Выравнивание"], [
    ["mean", "среднее трёх ветвей", "нет"],
    ["mean_mmd", "среднее трёх ветвей", "MMD"],
    ["mean_coral", "среднее трёх ветвей", "CORAL"],
    ["concat", "конкатенация", "нет"],
    ["concat_mmd", "конкатенация", "MMD"],
    ["concat_coral", "конкатенация", "CORAL"],
    ["gated", "обучаемые веса ветвей", "нет"],
    ["gated_coral", "обучаемые веса ветвей", "CORAL"],
], [3.3, 7.0, 5.4], 8.7)
paragraph("Контроли: визуальная, текстовая и числовая ветви отдельно; визуальная ветвь с MMD и с CORAL; zero-shot SigLIP 2. Всего обучено 13 × 3 × 2 = 78 голов на трёх фиксированных seed и двух наборах. Оптимизатор AdamW, 30 эпох, batch 64, learning rate 0,001, weight decay 0,0001, вес выравнивания 0,1. Нормировка каждого признака рассчитана только на source_train. В каждой тренировке эпоха выбиралась по Macro-F1 Product validation; схема — по среднему Macro-F1 этой же части между тремя seed. Clipart target_adapt использовался без истинных меток только в функции MMD/CORAL. Тестовые метки не входили в выбор или обучение этой версии кода.")
paragraph("Условие независимого подтверждения ограничено: целевые части Office-Home уже анализировались в прежнем варианте лабораторной. Хотя новый перечень схем и кодировщик закреплены до нынешнего пересчёта, выводы по Clipart, Art и Real World остаются исследовательскими. Статистические интервалы получены стратифицированным по классам bootstrap (2000 повторов); сравнения — парной перестановкой предсказаний (5000 повторов). Множественные сравнения не корректировались.")

heading("5 Результаты на Office-Home")
caption("Таблица 4 — Clipart test, ансамбль трёх seed для обучаемых схем")
order = ["visual", "text", "numeric", "visual_mmd", "visual_coral",
         "mean", "mean_mmd", "mean_coral", "concat", "concat_mmd",
         "concat_coral", "gated", "gated_coral", "zero_shot"]
table(["Схема", "Macro-F1", "95% ДИ", "Accuracy", "ECE"],
      [[key, fmt((r := lookup(schemes, dataset="officehome", scheme=key))["macro_f1"]),
        score_ci(r), fmt(r["accuracy"]), fmt(r["ece"])] for key in order],
      [3.5, 2.2, 4.3, 2.5, 2.5], 8.5)
o = lookup(chosen, dataset="officehome", split="test")
paragraph(f"По Product validation выбрана схема {selection['officehome']['fusion_scheme']}. На Clipart она получила Macro-F1 {fmt(o['fusion_macro_f1'])}; zero-shot SigLIP 2 — {fmt(o['unimodal_macro_f1'])}. Разница {fmt(o['difference'])}, 95% ДИ {ci(o)}, парный p={fmt(o['paired_permutation_two_sided_p'], 4)}. Положительного эффекта объединения нет. При этом визуальная обученная голова без дополнительных входов дала {fmt(lookup(schemes, dataset='officehome', scheme='visual')['macro_f1'])}, то есть сама сильная визуальная репрезентация уже хорошо переносится.")
caption("Рисунок 1 — Macro-F1 восьми схем объединения на двух целевых тестах")
doc.add_picture(str(RESULTS / "target_macro_f1.png"), width=Cm(15.7))

heading("6 Контроль выравнивания и второй набор")
caption("Таблица 5 — Парные сравнения на целевых тестах")
names = {"effect_of_mmd_visual": "MMD к визуальной модели",
         "effect_of_coral_visual": "CORAL к визуальной модели",
         "fusion_effect_given_mmd": "Слияние при MMD",
         "fusion_effect_given_coral": "Слияние при CORAL",
         "zero_shot_vs_trained_visual": "Zero-shot против visual"}
table(["Набор", "Сравнение", "Разница F1", "95% ДИ", "p"],
      [["Office-Home" if r["dataset"] == "officehome" else "Digits",
        names[r["comparison"]], fmt(r["difference"]), ci(r),
        fmt(r["paired_permutation_two_sided_p"], 4)] for r in matched],
      [2.2, 5.2, 2.0, 4.3, 2.0], 8.1)
d = lookup(chosen, dataset="digits", split="test")
same = lookup(matched, dataset="digits", comparison="fusion_effect_given_mmd")
paragraph(f"На Digits схема {selection['digits']['fusion_scheme']} получила Macro-F1 {fmt(d['fusion_macro_f1'])} против {fmt(d['unimodal_macro_f1'])} у обычной визуальной головы: прирост {fmt(d['difference'])}. Но визуальная голова с тем же MMD уже достигает {fmt(same['control_macro_f1'])}; добавление текста и чисел при фиксированном MMD меняет результат только на {fmt(same['difference'])}, ДИ {ci(same)}. Следовательно, основной выигрыш на этом искусственном сдвиге связан с выравниванием, а самостоятельный вклад дополнительных входов следует оценивать по этому контрольному сравнению.")
paragraph("Схема, выбранная на Product, дополнительно применена к Art и Real World. Они показывают перенос на другие стили изображения, но не являются более поздними временными периодами. Их показатели и парные интервалы хранятся в results/paired_comparisons.csv; из-за предыдущего просмотра этих частей они рассматриваются как дополнительная диагностика.")

doc.add_page_break()
heading("7 Вклад входов, пропуски и ошибки")
caption("Таблица 6 — Маскирование входа выбранной схемы на целевом тесте")
labels = {"full": "Все входы", "without_visual": "Без изображения",
          "without_text": "Без текста", "without_numeric": "Без чисел",
          "visual_only_at_inference": "Только изображение"}
table(["Условие", "Office-Home F1", "Digits F1"],
      [[labels[k], fmt(lookup(ablations, dataset="officehome", ablation=k)["macro_f1"]),
        fmt(lookup(ablations, dataset="digits", ablation=k)["macro_f1"])]
       for k in labels], [7.2, 4.2, 4.2], 8.8)
paragraph("Это проверка отсутствующего входа во время применения уже обученной головы: соответствующая ветвь исключается или заполняется нулями. На Office-Home удаление текста повышает Macro-F1 выбранной схемы с 0,634 до 0,699, а чисел — до 0,658; следовательно, эти входы в данной обученной конфигурации мешают. На Digits удаление текста снижает F1 с 0,832 до 0,731, удаление чисел — до 0,822. Маскирование не тождественно отдельному обучению модели без модальности; одномодальные контроли приведены в таблице 4 и сохранены для Digits в results/scheme_intervals.csv.")
caption("Таблица 7 — Повторное обучение выбранного семейства без одного входа")
table(["Исключённый вход", "Office-Home F1", "Digits F1"],
      [[labels[f"without_{key}"],
        fmt(lookup(retrained, dataset="officehome", ablation=f"without_{key}")["macro_f1"]),
        fmt(lookup(retrained, dataset="digits", ablation=f"without_{key}")["macro_f1"])]
       for key in ("visual", "text", "numeric")], [7.2, 4.2, 4.2], 8.8)
paragraph("При отдельном повторном обучении без текста Office-Home даёт 0,804 вместо 0,634 у полного варианта; без чисел — 0,637. На Digits без чисел получилось 0,842 вместо 0,832. Эти дополнительные абляции назначены после просмотра основных результатов и служат диагностикой, а не независимой проверкой гипотезы.")
off_err = [e for e in errors if e["dataset"] == "officehome"]
dig_err = [e for e in errors if e["dataset"] == "digits"]
paragraph(f"Выбранная схема ошиблась на {len(off_err)} из 130 рисунков Clipart и {len(dig_err)} из 405 целевых цифр. На {sum(e['baseline_correct'] == 'True' for e in off_err)} ошибках Clipart zero-shot baseline ответил правильно. Ошибок с уверенностью выше 0,8 соответственно {sum(float(e['fusion_confidence']) > .8 for e in off_err)} и {sum(float(e['fusion_confidence']) > .8 for e in dig_err)}. Для 65 классов два тестовых изображения на класс дают нестабильную по отдельным классам оценку; причинность единичной ошибки без ручной проверки не утверждается.")
sample = [e for e in off_err if e["baseline_correct"] == "True"][:3]
if sample:
    caption("Таблица 8 — Примеры ошибок объединения при правильном ответе zero-shot")
    table(["ID", "Класс", "Ответ схемы", "Автоматический текст"],
          [[e["id"], class_names[int(e["true_label"])],
            class_names[int(e["fusion_prediction"])], e["text"][:70]] for e in sample],
          [1.2, 3.3, 3.3, 7.9], 8)

heading("8 Качество, ресурсы и решение о применении")
vis = lookup(schemes, dataset="officehome", scheme="visual")
zs = lookup(schemes, dataset="officehome", scheme="zero_shot")
fus = lookup(schemes, dataset="officehome", scheme=selection["officehome"]["fusion_scheme"])
paragraph(f"Калибровка тоже важна: ECE на Clipart равна {fmt(zs['ece'])} для zero-shot, {fmt(vis['ece'])} для визуальной обученной головы и {fmt(fus['ece'])} для выбранной схемы. ECE зависит от разбиения вероятностей на корзины и при 130 объектах является грубой оценкой. Дополнительно рассчитаны accuracy, MCC, многоклассовая Brier score и NLL; полные значения находятся в results/scheme_intervals.csv.")
paragraph(f"Извлечение изображений и текстов SigLIP 2 на локальном CPU в пакетах по 16 заняло {fmt(extract['officehome']['seconds_total'], 1)} с для 1235 Office-Home объектов и {fmt(extract['digits']['seconds_total'], 1)} с для 1797 Digits. Это включает чтение изображения, вычисление чисел и кодирование уже готового текста, но не первоначальное получение меток Apple Vision. Файл весов SigLIP 2 занимает около 1,5 ГБ; после извлечения в проекте остаются компактные векторы и обучаемые головы. Время собственно голов измерено отдельно в results/metrics.csv. Денежных расходов внешнего API в этом запуске не было.")
paragraph("Для основной задачи Product → Clipart практический выбор — прямое zero-shot использование SigLIP 2, а не выбранная трёхвходовая схема. Для задачи сравнения методов выравнивания при фиксированном backbone следует отдельно сопоставлять visual, visual+MMD и visual+CORAL; их разницы малы относительно уже высокого качества признаков. На Digits выравнивание существенно, но это искусственный сдвиг и другая задача, поэтому его результат нельзя переносить на Office-Home.")

heading("9 Ограничения и воспроизводимость")
paragraph("Главные ограничения: небольшая и ранее просмотренная целевая часть Office-Home; два объекта на класс; множественные исследовательские сравнения; неизвестное пересечение данных предобучения SigLIP 2 с изображениями Office-Home; системно управляемая версия Apple Vision; текст и числовые признаки, зависящие от того же изображения; искусственный сдвиг Digits. Отсутствие настоящих временных рядов не позволяет заявить выполнение части задания о временной синхронизации и проверке на будущем периоде. Для такого вывода нужен иной набор с независимым текстовым и последовательным числовым потоком.")
paragraph("Воспроизведение обеспечивают зафиксированная ревизия SigLIP 2, requirements.txt, configs/protocol.json, исходные коды извлечения, обучения и анализа, data/schema.json, аудит, исключения, истории 78 основных и 18 абляционных обучений, сохранённые головы и вероятности предсказаний. Метки целевого домена не используются в обучении и выборе, что проверяется по коду. Числа демонстрационного финансового примера задания в результаты не переносились.")

heading("Список литературы")
for r in literature:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    p.add_run(f"{int(r['id'])}. {r['authors']}. {r['title']} // {r['venue']}. — {r['year']}. — URL: ")
    hyperlink(p, r["url"], r["url"])
    p.add_run(" (дата обращения: 22.10.2026). — Текст : электронный.")
    for run in p.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(9.5)
p = doc.add_paragraph()
p.add_run("Источники данных: ")
hyperlink(p, "Office-Home", "https://www.hemanthdv.org/officeHomeDataset.html")
p.add_run("; ")
hyperlink(p, "scikit-learn Digits", "https://scikit-learn.org/stable/modules/generated/sklearn.datasets.load_digits.html")
p.add_run(" (дата обращения: 22.10.2026).")

heading("Приложение А Состав числового входа")
paragraph("Каждый показатель вычислен по одному исходному изображению после приведения копии к размеру 128 × 128 пикселей; только отношение сторон берётся из исходных размеров. Значения не являются измерениями во времени.")
caption("Таблица А1 — Девятнадцать непосредственных характеристик изображения")
table(["Группа", "Число", "Определение"], [
    ["Средние R, G, B", 3, "среднее каждого цветового канала"],
    ["Разброс R, G, B", 3, "стандартное отклонение каждого канала"],
    ["Насыщенность и яркость", 4, "среднее и разброс S и V в HSV"],
    ["Серый тон", 4, "среднее, разброс и квантили 10% и 90%"],
    ["Тёмные и светлые пиксели", 2, "доли ниже 0,15 и выше 0,85"],
    ["Границы", 2, "средний градиент и доля сильных границ"],
    ["Геометрия", 1, "логарифм отношения ширины к высоте"],
], [4.6, 1.6, 9.5], 8.5)

doc.core_properties.title = "Лабораторная работа 2 Мультимодальные модели для адаптации визуального домена"
doc.core_properties.subject = "SigLIP 2, прямые числовые признаки, MMD и CORAL на Office-Home"
doc.core_properties.author = "А. Д. Якумова"
doc.save(OUTPUT)
print(OUTPUT)
